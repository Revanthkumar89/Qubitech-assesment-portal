
from flask import Flask, render_template, request, redirect, url_for, session, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
from pathlib import Path
import os
import hmac
import secrets

APP_DIR = Path(__file__).resolve().parent
DATA_FILE = APP_DIR / "candidate_responses.xlsx"

# Change this code before sharing the assessment link.
ACCESS_CODE = os.environ.get("QUBITECH_ACCESS_CODE", "QUBI2026")
ADMIN_PASSWORD = os.environ.get(
    "QUBITECH_ADMIN_PASSWORD", "ChangeThisPassword123!"
)

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY") or secrets.token_hex(32)

QUESTIONS = [
    {
        "id": "Q1",
        "title": "Prioritization and judgment",
        "prompt": "It is 10:00 AM. The Founder has an external stakeholder meeting at 11:00 AM, but the technical team has not sent the latest project status report. Another department head has not responded to a promised update, and two meetings are booked for 11:00 AM. What would you handle first, and why?",
    },
    {
        "id": "Q2",
        "title": "Professional communication",
        "prompt": "Draft a professional message to the technical team requesting the missing report before the 11:00 AM stakeholder meeting. The message should be clear, polite, time-bound, and action-oriented.",
    },
    {
        "id": "Q3",
        "title": "Coordination and escalation",
        "prompt": "The Founder has two important meetings at the same time. How would you resolve the conflict without making an unauthorized decision? What would you do if one organizer does not agree?",
    },
    {
        "id": "Q4",
        "title": "Reporting and follow-through",
        "prompt": "Prepare a brief end-of-day status update covering completed work, pending items, responsible owners, deadlines, and next steps. Use a simple table or structured bullets.",
    },
]


def ensure_workbook():
    if DATA_FILE.exists():
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"
    ws.append([
        "Submission ID", "Submitted At", "Candidate Name", "Email / Phone",
        "Q1 Prioritization and judgment",
        "Q2 Professional communication",
        "Q3 Coordination and escalation",
        "Q4 Reporting and follow-through"
    ])

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B56")
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    ws.row_dimensions[1].height = 36

    widths = [22, 22, 26, 28, 60, 60, 60, 60]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[
            __import__("openpyxl").utils.get_column_letter(i)
        ].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:H1"
    ws.sheet_view.showGridLines = False
    wb.save(DATA_FILE)


def append_submission(name, contact, answers):
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb["Responses"]

    ws.append([
        datetime.now().strftime("%Y%m%d%H%M%S%f"),
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        name.strip(),
        contact.strip(),
        answers["Q1"],
        answers["Q2"],
        answers["Q3"],
        answers["Q4"],
    ])

    for cell in ws[ws.max_row]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(DATA_FILE)


@app.before_request
def require_access_code():
    # The access-code page and admin route are handled separately.
    if request.endpoint in ("access", "admin"):
        return None

    if not session.get("assessment_access"):
        return redirect(url_for("access"))


@app.route("/access", methods=["GET", "POST"])
def access():
    error = ""

    if request.method == "POST":
        supplied_code = request.form.get("access_code", "")

        if hmac.compare_digest(supplied_code, ACCESS_CODE):
            session.clear()
            session["assessment_access"] = True
            return redirect(url_for("assessment"))

        error = "Incorrect access code. Please try again."

    return """
    <!doctype html>
    <html lang="en">
    <head>
      <meta charset="utf-8">
      <meta name="viewport" content="width=device-width, initial-scale=1">
      <title>Assessment Access | Qubitech</title>
      <style>
        body {
          font-family: Arial, sans-serif;
          background: #f4f6f8;
          color: #17212b;
          max-width: 460px;
          margin: 70px auto;
          padding: 20px;
        }
        .card {
          background: white;
          border: 1px solid #d9e0e6;
          border-radius: 12px;
          padding: 26px;
        }
        label {
          display: block;
          font-weight: bold;
          margin: 18px 0 7px;
        }
        input {
          box-sizing: border-box;
          width: 100%;
          padding: 12px;
          border: 1px solid #c8d1d9;
          border-radius: 7px;
          font-size: 16px;
        }
        button {
          width: 100%;
          margin-top: 18px;
          padding: 13px;
          background: #174b72;
          color: white;
          border: 0;
          border-radius: 7px;
          font-size: 16px;
          font-weight: bold;
          cursor: pointer;
        }
        .error { color: #a12626; }
      </style>
    </head>
    <body>
      <div class="card">
        <h2>Qubitech Smart Solutions Pvt. Ltd.</h2>
        <h3>Candidate Assessment</h3>
        <p>Enter the access code provided by the hiring team.</p>
        <p class="error">""" + error + """</p>
        <form method="post">
          <label for="access_code">Access code</label>
          <input id="access_code" name="access_code"
                 type="password" required autofocus>
          <button type="submit">Access assessment</button>
        </form>
      </div>
    </body>
    </html>
    """


@app.route("/", methods=["GET", "POST"])
def assessment():
    if request.method == "POST":
        name = request.form.get("candidate_name", "").strip()
        contact = request.form.get("contact", "").strip()

        if not name:
            return render_template(
                "assessment.html", questions=QUESTIONS
            )

        answers = {
            q["id"]: request.form.get(q["id"], "").strip()
            for q in QUESTIONS
        }

        if any(not answers[q["id"]] for q in QUESTIONS):
            return render_template(
                "assessment.html", questions=QUESTIONS
            )

        append_submission(name, contact, answers)
        return render_template("submitted.html", name=name)

    return render_template("assessment.html", questions=QUESTIONS)


@app.route("/admin")
def admin():
    supplied = request.args.get("key", "")

    if not hmac.compare_digest(supplied, ADMIN_PASSWORD):
        return "Not found", 404

    ensure_workbook()
    return send_file(
        DATA_FILE,
        as_attachment=True,
        download_name="Qubitech_Candidate_Responses.xlsx"
    )


if __name__ == "__main__":
    ensure_workbook()
    app.run(host="127.0.0.1", port=5000, debug=False)
